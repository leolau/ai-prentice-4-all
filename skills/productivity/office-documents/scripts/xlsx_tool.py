#!/usr/bin/env python3
"""Parse, read, update and create Excel (.xlsx) workbooks on the box (openpyxl).

Usage:
    python xlsx_tool.py info    BOOK.xlsx                          # sheets, dimensions, defined names
    python xlsx_tool.py dump    BOOK.xlsx [--sheet S] [--range A1:D20] [--formulas] [--json|--csv]
    python xlsx_tool.py create  OUT.xlsx --csv rows.csv [--sheet Name]
    python xlsx_tool.py create  OUT.xlsx --json rows.json [--sheet Name]   # [[...],[...]] or [{...},{...}]
    python xlsx_tool.py set     BOOK.xlsx [--sheet S] --cell B2=42 --cell C2="=A2*B2" [--out OUT.xlsx]
    python xlsx_tool.py append  BOOK.xlsx [--sheet S] --csv rows.csv | --json rows.json [--out OUT.xlsx]
    python xlsx_tool.py add-sheet BOOK.xlsx --sheet Name [--csv rows.csv] [--out OUT.xlsx]

Updates load the workbook with formulas preserved (not cached values) and
write it back in place unless --out is given. `dump` shows computed values
that were cached by the last save in Excel/Numbers; formulas entered here
have no cached value until the file is opened in a spreadsheet app (use
--formulas to see the formula text). Legacy .xls and Apple .numbers are not
supported — convert to .xlsx first.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date, datetime
from pathlib import Path


def _ensure() -> None:
    try:
        from tools.lazy_deps import ensure

        ensure("documents.openpyxl", prompt=False)
    except ImportError:
        pass


def _load(path: str, *, formulas: bool):
    _ensure()
    from openpyxl import load_workbook

    return load_workbook(path, data_only=not formulas)


def _sheet(wb, name: str | None):
    if name is None:
        return wb.active
    if name not in wb.sheetnames:
        raise SystemExit(f"no sheet {name!r}; sheets: {wb.sheetnames}")
    return wb[name]


def _coerce(raw: str):
    """Turn a CLI/CSV string into a typed cell value. Formulas stay strings."""
    if raw.startswith("="):
        return raw
    for conv in (int, float):
        try:
            return conv(raw)
        except ValueError:
            pass
    if raw.lower() in ("true", "false"):
        return raw.lower() == "true"
    if raw == "":
        return None
    return raw


def _json_default(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _rows_from_args(args: argparse.Namespace) -> list[list]:
    if getattr(args, "csv", None):
        with open(args.csv, newline="", encoding="utf-8") as fh:
            return [[_coerce(c) for c in row] for row in csv.reader(fh)]
    if getattr(args, "json", None):
        src = sys.stdin if args.json == "-" else open(args.json, encoding="utf-8")
        with src:
            data = json.load(src)
        if not isinstance(data, list):
            raise SystemExit("--json must be a list of rows (lists) or a list of objects")
        if data and isinstance(data[0], dict):
            headers = list(data[0].keys())
            return [headers] + [[row.get(h) for h in headers] for row in data]
        return [list(row) for row in data]
    return []


def cmd_info(args: argparse.Namespace) -> None:
    wb = _load(args.book, formulas=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append(
            {
                "name": ws.title,
                "dimensions": ws.dimensions,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "merged": [str(r) for r in ws.merged_cells.ranges][:50],
                "tables": list(ws.tables.keys()),
            }
        )
    print(
        json.dumps(
            {
                "file": args.book,
                "active": wb.active.title,
                "sheets": sheets,
                "defined_names": sorted(wb.defined_names.keys()),
            },
            indent=2,
            ensure_ascii=False,
        )
    )


def cmd_dump(args: argparse.Namespace) -> None:
    wb = _load(args.book, formulas=args.formulas)
    ws = _sheet(wb, args.sheet)
    cells = ws[args.range] if args.range else ws.iter_rows()
    rows = [[c.value for c in row] for row in cells]
    if args.json:
        print(json.dumps(rows, indent=2, ensure_ascii=False, default=_json_default))
    elif args.csv:
        w = csv.writer(sys.stdout)
        for row in rows:
            w.writerow(["" if v is None else v for v in row])
    else:
        print(f"# {ws.title} {args.range or ws.dimensions}")
        for i, row in enumerate(rows, start=1):
            print(f"{i:>4} | " + " | ".join("" if v is None else str(v) for v in row))


def cmd_create(args: argparse.Namespace) -> None:
    _ensure()
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet or "Sheet1"
    for row in _rows_from_args(args):
        ws.append(row)
    wb.save(args.output)
    _report(args.output)


def _target(args: argparse.Namespace) -> str:
    return args.out or args.book


def cmd_set(args: argparse.Namespace) -> None:
    wb = _load(args.book, formulas=True)
    ws = _sheet(wb, args.sheet)
    for item in args.cell:
        if "=" not in item:
            raise SystemExit(f"--cell expects REF=VALUE, got {item!r}")
        ref, value = item.split("=", 1)
        ws[ref] = _coerce(value)
    wb.save(_target(args))
    _report(_target(args))


def cmd_append(args: argparse.Namespace) -> None:
    wb = _load(args.book, formulas=True)
    ws = _sheet(wb, args.sheet)
    rows = _rows_from_args(args)
    if not rows:
        raise SystemExit("append needs --csv or --json")
    if args.skip_header:
        rows = rows[1:]
    for row in rows:
        ws.append(row)
    wb.save(_target(args))
    _report(_target(args))


def cmd_add_sheet(args: argparse.Namespace) -> None:
    wb = _load(args.book, formulas=True)
    if args.sheet in wb.sheetnames:
        raise SystemExit(f"sheet {args.sheet!r} already exists")
    ws = wb.create_sheet(args.sheet)
    for row in _rows_from_args(args):
        ws.append(row)
    wb.save(_target(args))
    _report(_target(args))


def _report(path: str) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    print(
        json.dumps(
            {
                "output": path,
                "bytes": Path(path).stat().st_size,
                "sheets": {ws.title: ws.dimensions for ws in wb.worksheets},
            }
        )
    )


def main(argv: list[str] | None = None) -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("info"); s.add_argument("book"); s.set_defaults(fn=cmd_info)

    s = sub.add_parser("dump"); s.add_argument("book"); s.add_argument("--sheet")
    s.add_argument("--range"); s.add_argument("--formulas", action="store_true")
    g = s.add_mutually_exclusive_group()
    g.add_argument("--json", action="store_true"); g.add_argument("--csv", action="store_true")
    s.set_defaults(fn=cmd_dump)

    s = sub.add_parser("create"); s.add_argument("output"); s.add_argument("--sheet")
    g = s.add_mutually_exclusive_group(); g.add_argument("--csv"); g.add_argument("--json")
    s.set_defaults(fn=cmd_create)

    s = sub.add_parser("set"); s.add_argument("book"); s.add_argument("--sheet")
    s.add_argument("--cell", action="append", default=[], required=True, help="REF=VALUE (repeatable)")
    s.add_argument("--out"); s.set_defaults(fn=cmd_set)

    s = sub.add_parser("append"); s.add_argument("book"); s.add_argument("--sheet")
    g = s.add_mutually_exclusive_group(); g.add_argument("--csv"); g.add_argument("--json")
    s.add_argument("--skip-header", action="store_true"); s.add_argument("--out")
    s.set_defaults(fn=cmd_append)

    s = sub.add_parser("add-sheet"); s.add_argument("book"); s.add_argument("--sheet", required=True)
    g = s.add_mutually_exclusive_group(); g.add_argument("--csv"); g.add_argument("--json")
    s.add_argument("--out"); s.set_defaults(fn=cmd_add_sheet)

    args = p.parse_args(argv)
    args.fn(args)


if __name__ == "__main__":
    main()
